"""xlsx generation for Competitive Calendar Extraction — mirrors
app/deleted_showtimes/batch_io.py's build_output_xlsx shape (openpyxl,
bold header row, frozen header, sized columns)."""

from __future__ import annotations

import io

import openpyxl
from openpyxl.styles import Font

from app.calendar_extract.prompt import FIELD_ORDER

COLUMN_WIDTHS = {
    "movie_title": 44, "release_date": 13, "studio": 12, "rating": 10,
    "release_day_tag": 14, "release_day_date": 15, "miscellaneous": 40,
}


def build_output_xlsx(rows: list[dict]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calendar"

    ws.append(list(FIELD_ORDER))
    for c in ws[1]:
        c.font = Font(bold=True)

    for row in rows:
        ws.append([row.get(f, "") for f in FIELD_ORDER])

    for i, field_name in enumerate(FIELD_ORDER, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = COLUMN_WIDTHS.get(field_name, 20)
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
